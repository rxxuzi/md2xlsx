import io
import sys
from datetime import datetime
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook

import md2xlsx


def build(tmp_path, text):
    src = tmp_path / "in.md"
    dst = tmp_path / "out.xlsx"
    src.write_text(text, encoding="utf-8")
    md2xlsx.convert(src, dst)
    return load_workbook(dst)


def cli(tmp_path, text, *args, capsys=None):
    """Run main() on `text`; returns (exit code, stdout, stderr)."""
    src = tmp_path / "in.md"
    src.write_text(text, encoding="utf-8")
    code = md2xlsx.main([str(src), *args])
    out = capsys.readouterr()
    return code, out.out, out.err


def test_sheets_and_pipe_table(tmp_path):
    wb = build(tmp_path, "#[A]\n| x | y |\n|---|---|\n| 1 | 2 |\n#[B]\nhi\n")
    assert wb.sheetnames == ["A", "B"]
    ws = wb["A"]
    assert ws["A1"].value == "x" and ws["A1"].font.bold
    assert ws["B2"].value == 2
    assert wb["B"]["A1"].value == "hi"


def test_csv_table_and_types(tmp_path):
    wb = build(tmp_path, "```table\nn,usd,jpy,pct,d\na,$1200,¥300,30%,2024-01-15\n```\n")
    ws = wb.active
    assert ws["B2"].value == 1200 and ws["B2"].number_format == "$#,##0"
    assert ws["C2"].number_format == "¥#,##0"
    assert ws["D2"].value == 0.3
    assert ws["E2"].value == datetime(2024, 1, 15)


def test_formula_and_link(tmp_path):
    wb = build(tmp_path, "| t | $$=SUM(B2:B3)$$ |\nsee [doc](https://x.y)\n")
    ws = wb.active
    assert ws["B1"].value == "=SUM(B2:B3)"
    assert ws["A2"].value == "see doc"
    assert ws["A2"].hyperlink.target == "https://x.y"


def test_decorators(tmp_path):
    wb = build(tmp_path, "| {bg:yellow}a | {fg:red}b | {dropdown:x,y}x | {comment:c}d |\n")
    ws = wb.active
    assert ws["A1"].fill.fgColor.rgb == "00FFFF00"
    assert ws["B1"].font.color.rgb == "00FF0000"
    assert len(ws.data_validations.dataValidation) == 1
    assert ws["D1"].comment.text == "c"


def test_directives(tmp_path):
    wb = build(
        tmp_path, "#[S]\n@freeze(1)\n@filter\n@style(A=5)\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    )
    ws = wb["S"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B2"
    assert ws.column_dimensions["A"].width == 5


def test_column_style(tmp_path):
    wb = build(
        tmp_path,
        "@style(a=12 yellow bold, B=#ABCDEF center italic)\n"
        "| h1 | h2 |\n|---|---|\n| x | {bg:red}{align:left}y |\n| z | w |\n",
    )
    ws = wb.active
    assert ws.column_dimensions["A"].width == 12
    assert ws["A2"].fill.fgColor.rgb == "00FFFF00" and ws["A2"].font.bold
    assert ws["B3"].fill.fgColor.rgb == "00ABCDEF" and ws["B3"].font.italic
    assert ws["B3"].alignment.horizontal == "center"
    # cell decorators win over the column style
    assert ws["B2"].fill.fgColor.rgb == "00FF0000"
    assert ws["B2"].alignment.horizontal == "left"
    # table headers are left alone entirely
    assert ws["A1"].fill.fgColor.rgb == "004472C4" and ws["A1"].font.color.rgb == "00FFFFFF"
    assert ws["B1"].alignment.horizontal == "center" and not ws["B1"].font.italic


def test_document_mode(tmp_path):
    wb = build(tmp_path, "@document\n# 1\n## 1.1\n### leaf\n")
    ws = wb.active
    assert ws["A1"].value == "1" or ws["A1"].value == 1
    assert ws["B2"].value == 1.1
    assert ws["C3"].value == "leaf" and not ws["C3"].font.bold
    assert ws["C3"].border.left.style == "thin"


def test_hr_and_cjk_width(tmp_path):
    wb = build(tmp_path, "キックオフ\n---\nx\n")
    ws = wb.active
    assert ws["A2"].value is None and ws["A3"].value == "x"
    assert ws.column_dimensions["A"].width >= 13


def test_document_body_text(tmp_path):
    md = "@document\n# 1. A\n## 1.1 B\nbody\n- one\n  - nested\n1. first\n# 2. C\nafter\n"
    wb = build(tmp_path, md)
    ws = wb.active
    assert ws["C3"].value == "body" and not ws["C3"].font.bold
    assert ws["C4"].value == "one" and ws["D5"].value == "nested"
    assert ws["C6"].value == "1. first"
    assert ws["B8"].value == "after"


def test_list_markers_default_mode(tmp_path):
    ws = build(tmp_path, "- x\n* y\n1. z\n  - indented\n").active
    assert [ws.cell(r, 1).value for r in range(1, 5)] == ["x", "y", "1. z", "indented"]
    assert ws["B4"].value is None


def test_error_format_rustc_style(tmp_path, capsys):
    code, out, err = cli(tmp_path, "#[S]\n@freez(1)\n| {bg:yelow}a |\n", capsys=capsys)
    assert code == 1 and out == "" and not (tmp_path / "in.xlsx").exists()
    assert "error: unknown directive `@freez`\n" in err
    assert "\n2 | @freez(1)\n  | ^^^^^^\n" in err
    assert f"--> {tmp_path / 'in.md'}:2:1" in err
    assert "= help: directives: @freeze(N)" in err
    # caret sits under the offending token, not the whole line
    assert "\n3 | | {bg:yelow}a |\n  |       ^^^^^\n" in err
    assert err.rstrip().endswith("error: aborting due to 2 previous errors")


def test_errors_cover_the_dialect(tmp_path):
    for bad in (
        "@style(A=huge)\n",
        "@style(A)\n",
        "@freeze(x)\n",
        "@filter(1)\n",
        "| {align:middle}x |\n",
        "| {fg:#12345}x |\n",
        "#[a/b]\n",
        "#[A]\nx\n#[A]\ny\n",
    ):
        with pytest.raises(ValueError):
            build(tmp_path, bad)


def test_warnings_do_not_block(tmp_path, capsys):
    md = "#[Doc]\n@document\n@style(A=4)\n# T\n| a | b |\n|---|---|\n| x \\| y | {zz:1}z |\n@filter\n"
    code, out, err = cli(tmp_path, md, capsys=capsys)
    assert code == 0 and (tmp_path / "in.xlsx").exists()
    for w in (
        "warning: column widths on a document sheet",
        "warning: table inside a document sheet",
        "warning: `\\|` does not escape `|`",
        "warning: unknown decorator `zz`",
        "warning: `@filter` here is plain text",
    ):
        assert w in err
    assert "^^^^^^^^^^^ column A holds the title" in err  # label after the carets
    assert err.rstrip().endswith("warning: 5 warnings emitted")
    assert out.startswith(str(tmp_path / "in.xlsx") + ": Doc")


def test_single_hash_warns_in_default_mode(tmp_path, capsys):
    code, _, err = cli(tmp_path, "# Title\n", capsys=capsys)
    assert code == 0 and "not a heading in a table sheet" in err


def test_cli_output_dump_verbose_quiet(tmp_path, capsys):
    md = "#[R]\n@freeze(1)\n| n | v |\n|---|---|\n| a | $1200 |\n| **t** | $$=SUM(B2:B2)$$ |\n"
    dst = tmp_path / "sub" / "r.xlsx"
    dst.parent.mkdir()
    code, out, err = cli(tmp_path, md, "-o", str(dst), "-d", "-v", capsys=capsys)
    assert code == 0 and err == "" and dst.exists()
    assert "#[R]  3x2  freeze=A2\n1 | *n | *v\n2 | a | $1,200\n3 | *t | =SUM(B2:B2)\n" in out
    assert f"{dst}: R\n  R: 3 rows, 2 cols, 1 formulas, @freeze(1)\n" in out
    code, out, _ = cli(tmp_path, md, "-o", str(dst), "-q", capsys=capsys)
    assert code == 0 and out == ""


def test_cli_stdin_and_missing_input(tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(sys, "stdin", io.TextIOWrapper(io.BytesIO(b"hi\n")))
    dst = tmp_path / "s.xlsx"
    assert md2xlsx.main(["-", "-o", str(dst)]) == 0 and dst.exists()
    capsys.readouterr()
    monkeypatch.setattr(sys, "stdin", io.TextIOWrapper(io.BytesIO(b"x\n")))
    assert md2xlsx.main(["-"]) == 1
    assert "`-o FILE` is required when reading from stdin" in capsys.readouterr().err
    assert md2xlsx.main([str(tmp_path / "nope.md")]) == 1
    assert "error: could not read `" in capsys.readouterr().err
