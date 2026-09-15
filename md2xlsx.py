#!/usr/bin/env python3
"""Convert flavored Markdown to xlsx. See docs/syntax.md for the dialect."""

import argparse
import csv
import io
import os
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

__version__ = "0.1.0"

FONT = "Arial"
BODY = Font(name=FONT, size=11)
BOLD = Font(name=FONT, size=11, bold=True)
H2 = Font(name=FONT, size=13, bold=True)
HEAD = Font(name=FONT, size=11, bold=True, color="FFFFFF")
CODE = Font(name="Consolas", size=11, color="C7254E")
LINK = Font(name=FONT, size=11, color="0563C1", underline="single")
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
ROW_BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))
GRID = Border(*(Side(style="thin", color="000000") for _ in range(4)))
CENTER = Alignment(horizontal="center", vertical="center")
MIDDLE = Alignment(vertical="center")

COLORS = {
    "red": "FF0000",
    "green": "00B050",
    "blue": "0070C0",
    "yellow": "FFFF00",
    "orange": "FFC000",
    "purple": "7030A0",
    "pink": "FF69B4",
    "gray": "808080",
    "grey": "808080",
    "white": "FFFFFF",
    "black": "000000",
    "lightblue": "BDD7EE",
    "lightgreen": "C6EFCE",
    "lightred": "FFC7CE",
    "lightyellow": "FFFFCC",
}
DIRECTIVES = ("freeze", "filter", "style", "document")
ALIGNS = ("left", "center", "right")
BAD_TITLE = set("[]:*?/\\")

RE_SHEET = re.compile(r"^#\[(.+?)\]\s*$")
RE_DIR = re.compile(r"^@(\w+)(?:\((.+?)\))?\s*$")
RE_FENCE = re.compile(r"^```table(?::(csv|tsv))?\s*$")
RE_HEAD = re.compile(r"^(#{1,6})\s+(.+)$")
RE_HR = re.compile(r"^\s*-{3,}\s*$")
RE_ITEM = re.compile(r"^(\s*)(?:[-*+]\s+|(\d+[.)])\s+)(.*)$")
RE_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
RE_FORMULA = re.compile(r"^\$\$(.+?)\$\$$")
RE_USD = re.compile(r"^\$([0-9,]+\.?[0-9]*)$")
RE_JPY = re.compile(r"^[¥￥]([0-9,]+\.?[0-9]*)$")
RE_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
RE_PCT = re.compile(r"^([0-9.]+)%$")
RE_LINK = re.compile(r"\[(.+?)\]\((.+?)\)")
RE_DECO = re.compile(r"\{(bg|fg|comment|dropdown|align):([^}]+)\}")
RE_ANY_DECO = re.compile(r"\{(\w+):[^}]*\}")
RE_HEX = re.compile(r"^#?[0-9a-fA-F]{6}$")
RE_BOLD = re.compile(r"\*\*(.+?)\*\*")
RE_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
RE_CODE = re.compile(r"`(.+?)`")


# ---------------------------------------------------------------- diagnostics


class Bad(Exception):
    """A dialect error in the line being processed. `token` is the offending text, so the
    renderer can put the caret under it; `help` is the one-line fix."""

    def __init__(self, msg, token=None, help=None):
        super().__init__(msg)
        self.token, self.help = token, help


ANSI = {"error": "1;31", "warning": "1;33", "blue": "1;34", "bold": "1"}


class Paint:
    def __init__(self, on):
        self.on = on

    def __call__(self, code, s):
        return f"\x1b[{ANSI[code]}m{s}\x1b[0m" if self.on else s


PLAIN = Paint(False)


def plural(n, s):
    return f"{n} {s}{'s' if n != 1 else ''}"


def ansi_ok(stream):
    if os.environ.get("NO_COLOR") or not hasattr(stream, "isatty") or not stream.isatty():
        return False
    if os.name != "nt":
        return True
    try:  # enable virtual-terminal sequences on the Windows console
        import ctypes

        k32 = ctypes.windll.kernel32
        h = k32.GetStdHandle(-12 if stream is sys.stderr else -11)
        mode = ctypes.c_uint32()
        return bool(
            k32.GetConsoleMode(h, ctypes.byref(mode)) and k32.SetConsoleMode(h, mode.value | 4)
        )
    except (AttributeError, OSError):
        return False


class Diags:
    """Collected errors and warnings, rendered rustc-style against the source lines."""

    def __init__(self, path="<input>", lines=()):
        self.path, self.lines, self.items = path, list(lines), []

    def error(self, msg, **kw):
        self.items.append(("error", msg, kw))

    def warn(self, msg, **kw):
        self.items.append(("warning", msg, kw))

    @property
    def errors(self):
        return sum(k == "error" for k, _, _ in self.items)

    @property
    def warnings(self):
        return sum(k == "warning" for k, _, _ in self.items)

    def render(self, paint=PLAIN):
        out = []
        for kind, msg, d in self.items:
            out.append(paint(kind, kind) + paint("bold", f": {msg}"))
            line = d.get("line")
            if line is None:
                out.extend(self.notes(d, "  ", paint))
            else:
                out.extend(self.snippet(kind, line, d, paint))
            out.append("")
        e, w = self.errors, self.warnings
        if e:
            tail = f"; {plural(w, 'warning')} emitted" if w else ""
            out.append(
                paint("error", "error")
                + paint("bold", f": aborting due to {plural(e, 'previous error')}{tail}")
            )
        elif w:
            out.append(
                paint("warning", "warning") + paint("bold", f": {plural(w, 'warning')} emitted")
            )
        return "\n".join(out)

    def snippet(self, kind, line, d, paint):
        src = self.lines[line] if line < len(self.lines) else ""
        token = d.get("token")
        at = src.find(token) if token else -1
        if at < 0:
            at, token = len(src) - len(src.lstrip()), src.strip()
        n = str(line + 1)
        pad = " " * len(n)
        bar = paint("blue", f"{pad} |")
        carets = "^" * max(width(token), 1) + (f" {d['label']}" if d.get("label") else "")
        out = [
            paint("blue", f"{pad}--> ") + f"{self.path}:{n}:{at + 1}",
            bar,
            paint("blue", f"{n} | ") + src,
            f"{bar} " + " " * width(src[:at]) + paint(kind, carets),
        ]
        notes = self.notes(d, f"{pad} ", paint)
        if notes:
            out.append(bar)
            out.extend(notes)
        return out

    @staticmethod
    def notes(d, indent, paint):
        return [
            f"{indent}{paint('blue', '=')} {paint('bold', key + ':')} {d[key]}"
            for key in ("help", "note")
            if d.get(key)
        ]


# ---------------------------------------------------------------- cell helpers


def color(name):
    n = name.strip().lower()
    if n in COLORS:
        return COLORS[n]
    if RE_HEX.match(n):
        return n.lstrip("#").upper()
    raise Bad(
        f"unknown color `{name.strip()}`",
        token=name.strip(),
        help="named colors: " + " ".join(COLORS) + "; or #RRGGBB",
    )


def num(s):
    for cast in (int, float):
        try:
            return cast(s)
        except ValueError:
            pass
    return s


def typed(text):
    """Return (value, number_format) for formula/currency/date/percent, else None."""
    if m := RE_FORMULA.match(text):
        return m.group(1), None
    if m := RE_USD.match(text):
        return float(m.group(1).replace(",", "")), "$#,##0"
    if m := RE_JPY.match(text):
        return float(m.group(1).replace(",", "")), "¥#,##0"
    if m := RE_DATE.match(text):
        return datetime(*map(int, m.groups())), "YYYY-MM-DD"
    if m := RE_PCT.match(text):
        return float(m.group(1)) / 100, "0.0%"
    return None


def decos(text):
    d = {}
    text = RE_DECO.sub(lambda m: d.__setitem__(m.group(1), m.group(2)) or "", text)
    return text.strip(), d


def refont(f, **kw):
    keep = ("name", "size", "bold", "italic", "underline", "color")
    return Font(**{**{k: getattr(f, k) for k in keep}, **kw})


def apply_decos(cell, d, sh):
    if "bg" in d:
        cell.fill = PatternFill("solid", fgColor=color(d["bg"]))
    if "fg" in d:
        cell.font = refont(cell.font, color=color(d["fg"]))
    if "comment" in d:
        cell.comment = Comment(d["comment"], "md2xlsx")
    if "dropdown" in d:
        opts = ",".join(o.strip() for o in d["dropdown"].split(","))
        dv = DataValidation(type="list", formula1=f'"{opts}"', showErrorMessage=True)
        sh.ws.add_data_validation(dv)
        dv.add(cell)
    if "align" in d:
        h = d["align"].strip().lower()
        if h not in ALIGNS:
            raise Bad(
                f"unknown alignment `{d['align'].strip()}`",
                token=d["align"].strip(),
                help="left, center or right",
            )
        cell.alignment = Alignment(horizontal=h, vertical="center")


def inline_font(text, base):
    bold = bool(RE_BOLD.search(text))
    italic = bool(RE_ITALIC.search(text))
    if RE_CODE.search(text) and not bold and not italic:
        return CODE
    if bold or italic:
        return Font(
            name=base.name, size=base.size, bold=bold or base.bold, italic=italic, color=base.color
        )
    return base


def strip_md(text):
    for rx in (RE_BOLD, RE_ITALIC, RE_CODE):
        text = rx.sub(r"\1", text)
    return text


def put(sh, col, raw, font=BODY):
    cell = sh.ws.cell(row=sh.row, column=col)
    text, d = decos(raw.strip())
    if m := RE_ANY_DECO.search(text):
        sh.warn(
            f"unknown decorator `{m.group(1)}`, written as text",
            token=m.group(0),
            help="decorators: bg, fg, align, comment, dropdown",
        )
    if not text:
        return cell
    if t := typed(text):
        cell.value, fmt = t
        cell.font = font
        if fmt:
            cell.number_format = fmt
    elif m := RE_LINK.search(text):
        cell.value = RE_LINK.sub(r"\1", text)
        cell.hyperlink = m.group(2)
        cell.font = LINK
    else:
        cell.value = num(strip_md(text))
        cell.font = inline_font(text, font)
    apply_decos(cell, d, sh)
    return cell


def width(s):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in s)


def auto_width(ws, skip=()):
    for cells in ws.columns:
        letter = get_column_letter(cells[0].column)
        if letter in skip:
            continue
        n = max((width(str(c.value)) for c in cells if c.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(n + 3, 8), 60)


def grid(ws):
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for c in row:
            c.border = GRID


def pipe_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def csv_rows(text, delim):
    if delim is None:
        delim = "\t" if text.count("\t") > text.count(",") else ","
    return list(csv.reader(io.StringIO(text.strip()), delimiter=delim))


def styles(spec):
    """@style(A=19 yellow bold, B=center) -> {"A": {"width": 19.0, "bg": "FFFF00", ...}}"""
    out = {}
    for pair in spec.split(","):
        col, eq, vals = pair.partition("=")
        col, pair = col.strip(), pair.strip()
        if not eq or not vals.strip():
            raise Bad(
                f"`{pair}` needs `=` and values", token=pair, help="COLUMN=values, e.g. A=6 center"
            )
        if not (col.isascii() and col.isalpha()):
            raise Bad(
                f"`{col}` is not a column letter",
                token=col,
                help="use the Excel column: A, B, …, AA",
            )
        st = {}
        for v in vals.split():
            lv = v.lower()
            if lv in ("bold", "italic"):
                st[lv] = True
            elif lv in ALIGNS:
                st["align"] = lv
            elif lv in COLORS or lv.startswith("#"):
                st["bg"] = color(v)
            else:
                try:
                    st["width"] = float(v)
                except ValueError:
                    raise Bad(
                        f"unrecognized value `{v}` in @style",
                        token=v,
                        help="a number (width), a color name or #RRGGBB, bold, italic, left, center or right",
                    ) from None
        out[col.upper()] = st
    return out


def style_cell(c, st):
    """Column style fills in only what the cell does not set itself."""
    if "bg" in st and c.fill.fill_type is None:
        c.fill = PatternFill("solid", fgColor=st["bg"])
    if "align" in st and c.alignment.horizontal is None:
        c.alignment = Alignment(horizontal=st["align"], vertical="center")
    kw = {k: True for k in ("bold", "italic") if st.get(k)}
    if kw:
        c.font = refont(c.font, **kw)


# ---------------------------------------------------------------- sheet


class Sheet:
    def __init__(self, wb, name, dirs, raw_dirs, diags):
        self.ws = wb.create_sheet(title=name)
        self.dirs, self.raw_dirs, self.diags = dirs, raw_dirs, diags
        self.doc = "document" in dirs
        self.at = 0  # source line index of what is being written, for diagnostics
        self.row = 0
        self.cols = 0
        self.formulas = 0
        self.heads = set()  # 1-based rows styled as table headers; @style skips them
        self.level = 0  # last heading level; document-mode body text goes one column deeper
        self.in_tbl = False
        self.headed = False

    def warn(self, msg, **kw):
        self.diags.warn(msg, line=self.at, **kw)

    def put(self, col, text, font=BODY):
        c = put(self, col, text, font)
        self.cols = max(self.cols, col)
        if c.data_type == "f":
            self.formulas += 1
        return c

    def blank(self):
        self.row += 1
        self.in_tbl = False

    def text(self, line):
        """Plain line or list item. `- x` drops the marker, `1. x` keeps it; in document
        mode the cell sits under the current heading, two spaces of indent = one column."""
        self.row += 1
        self.in_tbl = False
        col, text = 1, line.strip()
        if m := RE_ITEM.match(line):
            text = f"{m.group(2)} {m.group(3)}" if m.group(2) else m.group(3)
            if self.doc:
                col += len(m.group(1).expandtabs(2)) // 2
        if self.doc:
            col += self.level
        self.put(col, text)

    def heading(self, level, text):
        self.row += 1
        self.in_tbl = False
        if self.doc:
            self.level = level
            self.put(level, text, BOLD if level <= 2 else BODY)
            return
        c = self.put(1, text, H2 if level == 2 else BOLD)
        if c.alignment.horizontal is None:
            c.alignment = Alignment(horizontal="left")

    def head_row(self):
        self.heads.add(self.row)
        for c in self.ws[self.row]:
            if c.value is not None:
                c.font, c.fill, c.alignment = HEAD, HEAD_FILL, CENTER

    def table_start(self):
        self.in_tbl, self.headed = True, False
        if self.doc:
            self.warn(
                "table inside a document sheet",
                label="tables start at column A, on top of the outline",
                help="give the table its own sheet",
            )

    def table_row(self, cells, header=False):
        self.row += 1
        if header:
            self.heads.add(self.row)
        for i, v in enumerate(cells, start=1):
            c = self.put(i, v, HEAD if header else BODY)
            if header:
                c.alignment, c.fill = CENTER, HEAD_FILL
            elif c.alignment.horizontal is None:  # keep a {align:..} decorator
                c.alignment = MIDDLE
            if not self.doc:
                c.border = ROW_BORDER

    def pipe(self, line):
        if RE_SEP.match(line):
            if self.in_tbl and not self.headed:
                self.headed = True
                self.head_row()
            return
        if not self.in_tbl:
            self.table_start()
        if "\\|" in line:
            self.warn(
                "`\\|` does not escape `|`; the cell is split here",
                token="\\|",
                help="rephrase, or move the text to a document-sheet line where `|` is literal",
            )
        self.table_row(pipe_row(line))

    def csv(self, text, delim, start):
        self.at = start
        self.table_start()
        self.in_tbl = False
        for i, cells in enumerate(csv_rows(text, delim)):
            self.at = start + 1 + i
            self.table_row(cells, header=(i == 0))

    def finish(self):
        ws, d = self.ws, self.dirs
        st = d.get("style") or {}
        fixed = {col: s["width"] for col, s in st.items() if "width" in s}
        for col, v in fixed.items():
            ws.column_dimensions[col].width = v
        auto_width(ws, skip=fixed)
        if self.doc:
            grid(ws)
        for col, s in st.items():
            for c in ws[col]:
                if c.row not in self.heads:
                    style_cell(c, s)
        if "freeze" in d:
            ws.freeze_panes = f"A{d['freeze'] + 1}"
        if "filter" in d and self.cols and ws.max_row:
            ws.auto_filter.ref = f"A1:{get_column_letter(self.cols)}{ws.max_row}"


# ---------------------------------------------------------------- converter


class Conv:
    def __init__(self, diags=None):
        self.diags = diags or Diags()
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sh = None
        self.sheets = []
        self.named = {}  # sheet name -> line it was defined on
        self.fence = None
        self.fence_at = 0
        self.buf = []

    def sheet(self, name, dirs, raw_dirs, at):
        if self.sh:
            self.sh.finish()
        if bad := [c for c in name if c in BAD_TITLE]:
            self.diags.error(
                f"invalid character `{bad[0]}` in sheet name",
                line=at,
                token=name,
                help="Excel forbids [ ] : * ? / \\ in sheet names",
            )
            name = "".join(c for c in name if c not in BAD_TITLE) or "Sheet"
        if len(name) > 31:
            self.diags.warn(
                f"sheet name truncated to `{name[:31]}`",
                line=at,
                token=name,
                help="Excel allows 31 characters",
            )
            name = name[:31]
        if name in self.named:
            self.diags.error(
                f"duplicate sheet `{name}`",
                line=at,
                token=name,
                note=f"first defined at {self.diags.path}:{self.named[name] + 1}",
            )
        self.named.setdefault(name, at)
        self.sh = Sheet(self.wb, name, dirs, raw_dirs, self.diags)
        self.sheets.append(self.sh)

    def directive(self, name, arg):
        if name not in DIRECTIVES:
            raise Bad(
                f"unknown directive `@{name}`",
                token=f"@{name}",
                help="directives: @freeze(N), @filter, @style(...), @document",
            )
        if name == "freeze":
            if arg is None or not arg.strip().isdigit() or int(arg) < 1:
                raise Bad(
                    "`@freeze` needs a row count",
                    token=arg or "@freeze",
                    help="@freeze(1) freezes the header row",
                )
            return int(arg)
        if name == "style":
            if arg is None:
                raise Bad(
                    "`@style` needs column styles",
                    token="@style",
                    help="e.g. @style(A=6 center, B=30, D=lightgreen)",
                )
            return styles(arg)
        if arg is not None:
            raise Bad(f"`@{name}` takes no argument", token=arg, help=f"write `@{name}` alone")
        return True

    def read_dirs(self, lines, i):
        dirs, raw, where = {}, [], {}
        while i < len(lines):
            s = lines[i].strip()
            if s and not (m := RE_DIR.match(s)):
                break
            if s:
                name, arg = m.group(1).lower(), m.group(2)
                try:
                    dirs[name] = self.directive(name, arg)
                    raw.append(s)
                    where[name] = i
                except Bad as e:
                    self.diags.error(str(e), line=i, token=e.token, help=e.help)
            i += 1
        if "document" in dirs and any("width" in s for s in (dirs.get("style") or {}).values()):
            self.diags.warn(
                "column widths on a document sheet",
                line=where["style"],
                label="column A holds the title; auto-width fits the outline",
                help="drop the widths, or make this a table sheet",
            )
        return dirs, raw, i

    def run(self, lines):
        i = 0
        while i < len(lines):
            i = self.step(lines, i)
        if self.sh:
            self.sh.finish()
        if not self.wb.sheetnames:
            self.wb.create_sheet("Sheet1")
        return self.wb

    def step(self, lines, i):
        line = lines[i]
        if self.fence is not None:
            if line.strip() == "```":
                text, delim = "\n".join(self.buf), self.fence or None
                self.guard(lambda: self.sh.csv(text, delim, self.fence_at))
                self.fence, self.buf = None, []
            else:
                self.buf.append(line)
            return i + 1
        if m := RE_SHEET.match(line):
            dirs, raw, j = self.read_dirs(lines, i + 1)
            self.sheet(m.group(1).strip(), dirs, raw, i)
            return j
        if self.sh is None:
            if RE_DIR.match(line.strip()):
                dirs, raw, j = self.read_dirs(lines, i)
                self.sheet("Sheet1", dirs, raw, i)
                return j
            self.sheet("Sheet1", {}, [], i)
        self.sh.at = i
        self.guard(lambda: self.body(line))
        return i + 1

    def guard(self, fn):
        try:
            fn()
        except Bad as e:
            self.diags.error(str(e), line=self.sh.at, token=e.token, help=e.help)

    def body(self, line):
        sh, s = self.sh, line.strip()
        if m := RE_FENCE.match(s):
            self.fence, self.fence_at, self.buf = m.group(1) or "", sh.at, []
        elif RE_HR.match(line):
            sh.blank()
        elif (m := RE_HEAD.match(line)) and (sh.doc or len(m.group(1)) >= 2):
            sh.heading(len(m.group(1)), m.group(2).strip())
        elif "|" in line and not line.lstrip().startswith("#"):
            sh.pipe(line)
        elif not s:
            sh.in_tbl = False
        else:
            if (m := RE_DIR.match(s)) and m.group(1).lower() in DIRECTIVES:
                sh.warn(
                    f"`@{m.group(1)}` here is plain text",
                    label="directives are only recognised directly under #[Sheet]",
                    help="move it up, right below the #[...] line",
                )
            elif RE_HEAD.match(line):
                sh.warn(
                    "`# ` is not a heading in a table sheet; written as text",
                    help="use `## ` here, or `@document` for an outline",
                )
            sh.text(line)


def run(text, path="<input>"):
    """Convert Markdown text. Returns (workbook, converter, diagnostics); the workbook is
    only meaningful when diagnostics.errors == 0."""
    lines = text.splitlines()
    diags = Diags(path, lines)
    conv = Conv(diags)
    return conv.run(lines), conv, diags


def convert(src, dst):
    """File-to-file helper for library use and tests. Raises ValueError on dialect errors."""
    wb, _, diags = run(Path(src).read_text(encoding="utf-8"), str(src))
    if diags.errors:
        raise ValueError(diags.render())
    wb.save(dst)
    return wb.sheetnames


# ---------------------------------------------------------------- cli


def show(c):
    """Cell value as the user would see it in Excel, for --dump."""
    v, fmt = c.value, c.number_format
    if v is None:
        return ""
    if c.data_type == "f":
        return v
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and fmt.startswith("$"):
        return f"${v:,.0f}"
    if isinstance(v, float) and fmt.startswith("¥"):
        return f"¥{v:,.0f}"
    if isinstance(v, float) and fmt.endswith("%"):
        return f"{v * 100:g}%"
    if isinstance(v, float):
        return f"{v:g}"
    return str(v)


def dump(wb):
    out = []
    for ws in wb.worksheets:
        head = [f"#[{ws.title}]", f"{ws.max_row}x{ws.max_column}"]
        if ws.freeze_panes:
            head.append(f"freeze={ws.freeze_panes}")
        if ws.auto_filter.ref:
            head.append(f"filter={ws.auto_filter.ref}")
        out.append("  ".join(head))
        w = len(str(ws.max_row))
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            cells = [("*" if c.font.bold and c.value is not None else "") + show(c) for c in row]
            while cells and not cells[-1]:
                cells.pop()
            out.append(f"{row[0].row:>{w}} | " + " | ".join(cells))
        out.append("")
    return "\n".join(out).rstrip()


def main(argv=None):
    p = argparse.ArgumentParser(
        prog="md2xlsx",
        description="Convert flavored Markdown to a multi-sheet .xlsx workbook.",
        epilog="Dialect reference: docs/syntax.md. Errors and warnings go to stderr.",
    )
    p.add_argument("input", help="Markdown file, or - for stdin")
    p.add_argument(
        "-o", "--output", metavar="FILE", help="output .xlsx (default: INPUT with .xlsx)"
    )
    p.add_argument("-d", "--dump", action="store_true", help="print each sheet as a text grid")
    g = p.add_mutually_exclusive_group()
    g.add_argument("-q", "--quiet", action="store_true", help="no summary line")
    g.add_argument("-v", "--verbose", action="store_true", help="per-sheet summary")
    p.add_argument("--version", action="version", version=f"%(prog)s {__version__}")
    a = p.parse_args(argv)

    for s in (sys.stdout, sys.stderr):
        if hasattr(s, "reconfigure"):
            s.reconfigure(errors="replace")
    paint = Paint(ansi_ok(sys.stderr))

    def fail(msg, **kw):
        d = Diags()
        d.error(msg, **kw)
        print(d.render(paint), file=sys.stderr)
        return 1

    if a.input == "-":
        if not a.output:
            return fail("`-o FILE` is required when reading from stdin")
        text, path = sys.stdin.buffer.read().decode("utf-8", errors="replace"), "<stdin>"
    else:
        path = a.input
        try:
            text = Path(path).read_text(encoding="utf-8")
        except OSError as e:
            return fail(f"could not read `{path}`: {e.strerror or e}")
        except UnicodeDecodeError:
            return fail(f"could not read `{path}`: not valid UTF-8", help="save the file as UTF-8")
    dst = a.output or str(Path(a.input).with_suffix(".xlsx"))

    wb, conv, diags = run(text, path)
    if diags.items:
        print(diags.render(paint), file=sys.stderr)
    if diags.errors:
        return 1
    if a.dump:
        print(dump(wb))
    try:
        wb.save(dst)
    except PermissionError:
        return fail(f"could not write `{dst}`: permission denied", help="is it open in Excel?")
    except OSError as e:
        return fail(f"could not write `{dst}`: {e.strerror or e}")
    if not a.quiet:
        print(f"{dst}: {', '.join(wb.sheetnames)}")
    if a.verbose:
        for sh in conv.sheets:
            dirs = " ".join(sh.raw_dirs) or "-"
            print(f"  {sh.ws.title}: {sh.row} rows, {sh.cols} cols, {sh.formulas} formulas, {dirs}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
